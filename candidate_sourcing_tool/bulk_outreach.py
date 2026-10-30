from __future__ import annotations
import base64, hashlib, json, os, re, sqlite3
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import urlparse
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import RedirectResponse
from pydantic import BaseModel, Field

router = APIRouter(prefix="/api/bulk-outreach", tags=["Bulk CV Outreach"])
APP_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("CANDIDATE_SOURCING_DB_PATH", APP_DIR / "candidate_sourcing.db"))
TOKEN_PATH = Path(os.getenv("GMAIL_TOKEN_PATH", APP_DIR / ".gmail_token.json"))
CLIENT_SECRETS_PATH = Path(os.getenv("GOOGLE_OAUTH_CLIENT_SECRETS", APP_DIR / "credentials.json"))
REDIRECT_URI = os.getenv("GOOGLE_OAUTH_REDIRECT_URI", "http://127.0.0.1:8000/api/bulk-outreach/gmail/callback")
SCOPES = ["https://www.googleapis.com/auth/gmail.compose"]
MAX_FILES, MAX_FILE_BYTES = 1000, 15 * 1024 * 1024
PLATFORMS = {"Turing", "Mercor", "micro1"}

def now():
    return datetime.now(timezone.utc).isoformat()

def connect():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DB_PATH, timeout=30)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA busy_timeout=30000")
    c.execute("PRAGMA journal_mode=WAL")
    return c

def init_tables():
    with connect() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS outreach_candidates (
          id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT NOT NULL UNIQUE, name TEXT NOT NULL DEFAULT '',
          skills_json TEXT NOT NULL DEFAULT '[]', experience_summary TEXT NOT NULL DEFAULT '',
          file_hashes_json TEXT NOT NULL DEFAULT '[]', source_filenames_json TEXT NOT NULL DEFAULT '[]',
          contact_basis_confirmed INTEGER NOT NULL DEFAULT 0, opted_out INTEGER NOT NULL DEFAULT 0,
          imported_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS outreach_opportunities (
          id INTEGER PRIMARY KEY AUTOINCREMENT, platform TEXT NOT NULL, role_title TEXT NOT NULL,
          requirements TEXT NOT NULL, referral_url TEXT NOT NULL, link_verified_by_user INTEGER NOT NULL DEFAULT 0,
          created_at TEXT NOT NULL, UNIQUE(platform,role_title,referral_url));
        CREATE TABLE IF NOT EXISTS outreach_drafts (
          id INTEGER PRIMARY KEY AUTOINCREMENT, candidate_id INTEGER NOT NULL, opportunity_id INTEGER NOT NULL,
          gmail_draft_id TEXT NOT NULL DEFAULT '', subject TEXT NOT NULL, body TEXT NOT NULL,
          match_score INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'LOCAL_DRAFT', created_at TEXT NOT NULL,
          UNIQUE(candidate_id,opportunity_id));
        CREATE TABLE IF NOT EXISTS outreach_batches (
          id INTEGER PRIMARY KEY AUTOINCREMENT, uploaded_count INTEGER NOT NULL, parsed_count INTEGER NOT NULL,
          unique_count INTEGER NOT NULL, duplicate_count INTEGER NOT NULL, error_count INTEGER NOT NULL, created_at TEXT NOT NULL);
        """)
init_tables()

class OpportunityCreate(BaseModel):
    platform: str
    role_title: str = Field(min_length=2, max_length=180)
    requirements: str = Field(min_length=2, max_length=10000)
    referral_url: str = Field(min_length=10, max_length=2000)
    link_verified_by_user: bool = False
class ConsentUpdate(BaseModel):
    contact_basis_confirmed: bool
class OptOutUpdate(BaseModel):
    opted_out: bool = True
class DraftCreate(BaseModel):
    candidate_id: int
    opportunity_id: int

EMAIL_RE = re.compile(r"(?i)(?<![\w.+-])[A-Z0-9.!#$%&'*+/=?^_{}|~-]+@[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?(?:\.[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?)+")
SKILLS = ["python","java","kotlin","swift","javascript","typescript","react","node.js","fastapi","django","flask","sql","postgresql","mysql","mongodb","aws","azure","gcp","docker","kubernetes","machine learning","deep learning","pytorch","tensorflow","llm","prompt engineering","ai evaluation","data annotation","data analysis","pandas","numpy","power bi","tableau","android","ios","flutter","react native","xgboost","nlp","computer vision","generative ai","data science"]

def extract_text(filename, data):
    if Path(filename).suffix.lower() == ".pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as pdf:
            text = "\n".join(p.get_text("text") for p in pdf)
        if not text.strip(): raise ValueError("PDF has no extractable text; scanned PDFs need local OCR.")
        return text
    if Path(filename).suffix.lower() == ".docx":
        import io
        from docx import Document
        d = Document(io.BytesIO(data))
        pieces = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows: pieces.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(pieces)
        if not text.strip(): raise ValueError("DOCX contains no extractable text.")
        return text
    raise ValueError("Only PDF and DOCX files are supported.")

def parse_candidate(text, filename):
    emails = list(dict.fromkeys(x.strip(".,;:()[]{}<>").lower() for x in EMAIL_RE.findall(text)))
    if not emails: return None
    lines = [re.sub(r"\s+"," ",x).strip() for x in text.splitlines() if x.strip()]
    name = ""
    for line in lines[:12]:
        cleaned = re.sub(r"[^A-Za-z .'-]", "", line).strip()
        words = cleaned.split()
        if 2 <= len(words) <= 5 and not any(x in cleaned.lower() for x in ["resume","curriculum vitae","profile","email","phone"]) and not any(s in cleaned.lower() for s in SKILLS):
            name = cleaned; break
    exp = re.search(r"(?i)(\d{1,2}\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:experience|exp))", text)
    return {"email": emails[0], "name": name or Path(filename).stem.replace("_"," ").replace("-"," "), "skills": sorted({s for s in SKILLS if s in text.lower()}), "experience": exp.group(1) if exp else ""}

def score(skills, requirements):
    req = requirements.lower()
    terms = [s for s in SKILLS if s in req]
    if not terms: terms = [w for w in re.findall(r"[a-z][a-z0-9+#.-]{2,}", req) if w not in {"and","the","with","years","experience","required"}]
    matched = [x for x in terms if x in skills]
    missing = [x for x in terms if x not in skills]
    return (min(100, round(100*len(matched)/max(1,len(terms)))), matched[:12], missing[:12]) if terms else (0,[],[])

def gmail_creds():
    try:
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
    except ImportError: return None, "Install Gmail dependencies from requirements.txt first."
    if not TOKEN_PATH.exists(): return None, "Gmail is not connected. Add Google OAuth credentials and connect Gmail."
    try:
        creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request()); TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
        if not creds.valid: return None, "Gmail credentials are invalid. Reconnect Gmail."
        return creds, ""
    except Exception as e: return None, f"Could not load Gmail credentials: {e}"

def gmail_service():
    creds, error = gmail_creds()
    if error: raise HTTPException(400, error)
    try:
        from googleapiclient.discovery import build
        return build("gmail","v1",credentials=creds,cache_discovery=False)
    except ImportError: raise HTTPException(400,"Install Gmail dependencies from requirements.txt first.")

@router.get("/status")
def status():
    with connect() as c:
        row = c.execute("SELECT (SELECT COUNT(*) FROM outreach_candidates) candidates,(SELECT COUNT(*) FROM outreach_opportunities) opportunities,(SELECT COUNT(*) FROM outreach_drafts) drafts,(SELECT COUNT(*) FROM outreach_candidates WHERE opted_out=1) opted_out").fetchone()
    _, err = gmail_creds() if TOKEN_PATH.exists() else (None,"not connected")
    return {**dict(row), "gmail_connected": not bool(err), "max_files_per_batch": MAX_FILES}

@router.post("/import")
async def import_batch(files: list[UploadFile] = File(...), contact_basis_confirmed: bool = Form(False)):
    if not contact_basis_confirmed: raise HTTPException(400,"Confirm you have an appropriate basis to contact these candidates about relevant opportunities.")
    if not files or len(files)>MAX_FILES: raise HTTPException(400,f"Upload between 1 and {MAX_FILES} files.")
    imported=duplicates=parsed=0; errors=[]
    with connect() as c:
        for f in files:
            filename = Path(f.filename or "unnamed").name
            if Path(filename).suffix.lower() not in {".pdf",".docx"}:
                errors.append({"filename":filename,"error":"Unsupported format; use PDF or DOCX."}); continue
            try:
                data=await f.read()
                if len(data)>MAX_FILE_BYTES: raise ValueError("File exceeds 15 MB limit.")
                text=extract_text(filename,data); parsed+=1; rec=parse_candidate(text,filename)
                if not rec: raise ValueError("No valid email address found.")
                h=hashlib.sha256(data).hexdigest()
                old=c.execute("SELECT * FROM outreach_candidates WHERE email=?",(rec["email"],)).fetchone()
                if old:
                    skills=sorted(set(json.loads(old["skills_json"] or "[]"))|set(rec["skills"]))
                    hashes=sorted(set(json.loads(old["file_hashes_json"] or "[]"))|{h})
                    names=sorted(set(json.loads(old["source_filenames_json"] or "[]"))|{filename})
                    c.execute("UPDATE outreach_candidates SET skills_json=?,file_hashes_json=?,source_filenames_json=?,contact_basis_confirmed=1,updated_at=? WHERE id=?",(json.dumps(skills),json.dumps(hashes),json.dumps(names),now(),old["id"]))
                    duplicates+=1
                else:
                    c.execute("INSERT INTO outreach_candidates(email,name,skills_json,experience_summary,file_hashes_json,source_filenames_json,contact_basis_confirmed,opted_out,imported_at,updated_at) VALUES(?,?,?,?,?,?,1,0,?,?)",(rec["email"],rec["name"],json.dumps(rec["skills"]),rec["experience"],json.dumps([h]),json.dumps([filename]),now(),now()))
                    imported+=1
            except Exception as e: errors.append({"filename":filename,"error":str(e)[:300]})
        c.execute("INSERT INTO outreach_batches(uploaded_count,parsed_count,unique_count,duplicate_count,error_count,created_at) VALUES(?,?,?,?,?,?)",(len(files),parsed,imported,duplicates,len(errors),now()))
    return {"uploaded_count":len(files),"parsed_count":parsed,"new_unique_candidates":imported,"duplicates_merged":duplicates,"errors":errors[:100],"error_count":len(errors),"message":"Processed locally. No emails were sent and import did not create drafts."}

@router.get("/candidates")
def candidates():
    with connect() as c: rows=c.execute("SELECT * FROM outreach_candidates ORDER BY updated_at DESC").fetchall()
    out=[]
    for r in rows:
        x=dict(r); x["skills"]=json.loads(x.pop("skills_json") or "[]"); x["source_filenames"]=json.loads(x.pop("source_filenames_json") or "[]"); x.pop("file_hashes_json",None); out.append(x)
    return out

@router.post("/candidates/{candidate_id}/contact-basis")
def set_basis(candidate_id:int,payload:ConsentUpdate):
    with connect() as c:
        cur=c.execute("UPDATE outreach_candidates SET contact_basis_confirmed=?,updated_at=? WHERE id=?",(int(payload.contact_basis_confirmed),now(),candidate_id))
        if not cur.rowcount: raise HTTPException(404,"Candidate not found.")
    return {"candidate_id":candidate_id,"contact_basis_confirmed":payload.contact_basis_confirmed}

@router.post("/candidates/{candidate_id}/opt-out")
def set_opt_out(candidate_id:int,payload:OptOutUpdate):
    with connect() as c:
        cur=c.execute("UPDATE outreach_candidates SET opted_out=?,updated_at=? WHERE id=?",(int(payload.opted_out),now(),candidate_id))
        if not cur.rowcount: raise HTTPException(404,"Candidate not found.")
    return {"candidate_id":candidate_id,"opted_out":payload.opted_out}

@router.get("/opportunities")
def opportunities():
    with connect() as c: return [dict(x) for x in c.execute("SELECT * FROM outreach_opportunities ORDER BY created_at DESC").fetchall()]

@router.post("/opportunities")
def add_opportunity(p:OpportunityCreate):
    if p.platform not in PLATFORMS: raise HTTPException(400,"Platform must be Turing, Mercor or micro1.")
    u=urlparse(p.referral_url.strip())
    if u.scheme!="https" or not u.netloc: raise HTTPException(400,"Referral link must be a valid HTTPS URL.")
    if not p.link_verified_by_user: raise HTTPException(400,"Verify the referral URL and platform terms before saving.")
    with connect() as c:
        try:
            cur=c.execute("INSERT INTO outreach_opportunities(platform,role_title,requirements,referral_url,link_verified_by_user,created_at) VALUES(?,?,?,?,1,?)",(p.platform,p.role_title.strip(),p.requirements.strip(),p.referral_url.strip(),now()))
            row=c.execute("SELECT * FROM outreach_opportunities WHERE id=?",(cur.lastrowid,)).fetchone()
        except sqlite3.IntegrityError:
            row=c.execute("SELECT * FROM outreach_opportunities WHERE platform=? AND role_title=? AND referral_url=?",(p.platform,p.role_title.strip(),p.referral_url.strip())).fetchone()
    return dict(row)

@router.get("/matches")
def matches(min_score:int=1):
    with connect() as c:
        cs=c.execute("SELECT * FROM outreach_candidates WHERE opted_out=0").fetchall()
        os_=c.execute("SELECT * FROM outreach_opportunities WHERE link_verified_by_user=1").fetchall()
        drafted={(x["candidate_id"],x["opportunity_id"]) for x in c.execute("SELECT candidate_id,opportunity_id FROM outreach_drafts").fetchall()}
    out=[]
    for cand in cs:
        skills=json.loads(cand["skills_json"] or "[]")
        for opp in os_:
            val,matched,missing=score(skills,opp["requirements"])
            if val<min_score: continue
            out.append({"candidate_id":cand["id"],"name":cand["name"],"email":cand["email"],"skills":skills,"contact_basis_confirmed":bool(cand["contact_basis_confirmed"]),"opportunity_id":opp["id"],"platform":opp["platform"],"role_title":opp["role_title"],"referral_url":opp["referral_url"],"match_score":val,"matched_skills":matched,"missing_skills":missing,"draft_exists":(cand["id"],opp["id"]) in drafted,"eligible_for_draft":bool(cand["contact_basis_confirmed"]) and not bool(cand["opted_out"]) and val>0})
    return sorted(out,key=lambda x:(x["match_score"],x["name"].lower()),reverse=True)

def email_content(name,platform,role,url,matched):
    first=name.split()[0] if name.strip() else "there"
    subject=f"Potential {role} opportunity via {platform}"
    skill=f"I noticed your background includes {', '.join(matched[:4])}, which may be relevant." if matched else "Your profile may be relevant."
    body=(f"Hi {first},\n\nI'm sharing a potential opportunity through {platform} for a {role} role. {skill}\n\n"
          "This is an independent referral, not an official shortlist, job offer, or guarantee of selection. "
          "The platform determines eligibility and next steps.\n\n"
          f"Review the opportunity and referral link here:\n{url}\n\n"
          "If this isn't relevant or you don't want further messages from me, please let me know.\n\nBest regards,\nRecruitment Team")
    return subject,body

@router.post("/drafts")
def create_draft(p:DraftCreate):
    with connect() as c:
        cand=c.execute("SELECT * FROM outreach_candidates WHERE id=?",(p.candidate_id,)).fetchone()
        opp=c.execute("SELECT * FROM outreach_opportunities WHERE id=?",(p.opportunity_id,)).fetchone()
        if not cand or not opp: raise HTTPException(404,"Candidate or opportunity not found.")
        if cand["opted_out"]: raise HTTPException(400,"Candidate opted out; draft blocked.")
        if not cand["contact_basis_confirmed"]: raise HTTPException(400,"Confirm appropriate contact basis before drafting.")
        if not opp["link_verified_by_user"]: raise HTTPException(400,"Referral link is not verified.")
        old=c.execute("SELECT * FROM outreach_drafts WHERE candidate_id=? AND opportunity_id=?",(p.candidate_id,p.opportunity_id)).fetchone()
        if old: return {"id":old["id"],"status":old["status"],"gmail_draft_id":old["gmail_draft_id"],"message":"Draft already exists."}
        skills=json.loads(cand["skills_json"] or "[]"); val,matched,_=score(skills,opp["requirements"])
        if val<1: raise HTTPException(400,"No evidence-based skill match; review candidate first.")
        subject,body=email_content(cand["name"],opp["platform"],opp["role_title"],opp["referral_url"],matched)
        msg=EmailMessage(); msg["To"]=cand["email"]; msg["Subject"]=subject; msg.set_content(body)
        encoded=base64.urlsafe_b64encode(msg.as_bytes()).decode("ascii")
        service=gmail_service()
        try: result=service.users().drafts().create(userId="me",body={"message":{"raw":encoded}}).execute()
        except Exception as e: raise HTTPException(502,f"Gmail draft creation failed: {e}") from e
        cur=c.execute("INSERT INTO outreach_drafts(candidate_id,opportunity_id,gmail_draft_id,subject,body,match_score,status,created_at) VALUES(?,?,?,?,?,?,'GMAIL_DRAFT',?)",(cand["id"],opp["id"],result.get("id",""),subject,body,val,now()))
    return {"id":cur.lastrowid,"status":"GMAIL_DRAFT","gmail_draft_id":result.get("id",""),"subject":subject,"match_score":val,"message":"Draft created in Gmail. It has not been sent."}

@router.get("/drafts")
def drafts():
    with connect() as c:
        rows=c.execute("SELECT d.*,c.name,c.email,o.platform,o.role_title,o.referral_url FROM outreach_drafts d JOIN outreach_candidates c ON c.id=d.candidate_id JOIN outreach_opportunities o ON o.id=d.opportunity_id ORDER BY d.created_at DESC").fetchall()
    return [dict(x) for x in rows]

@router.get("/gmail/connect")
def gmail_connect():
    if not CLIENT_SECRETS_PATH.exists(): raise HTTPException(400,f"Google OAuth client JSON not found at {CLIENT_SECRETS_PATH}. Configure OAuth and save as credentials.json.")
    try:
        from google_auth_oauthlib.flow import Flow
        flow=Flow.from_client_secrets_file(str(CLIENT_SECRETS_PATH),scopes=SCOPES); flow.redirect_uri=REDIRECT_URI
        url,state=flow.authorization_url(access_type="offline",include_granted_scopes="true",prompt="consent")
        (APP_DIR/".gmail_oauth_state").write_text(state,encoding="utf-8")
        return {"authorization_url":url,"message":"Open the URL and authorize Gmail draft access."}
    except ImportError: raise HTTPException(400,"Install Gmail dependencies from requirements.txt first.")

@router.get("/gmail/callback")
def gmail_callback(code:str="",state:str="",error:str=""):
    if error: raise HTTPException(400,f"Google authorization failed: {error}")
    sp=APP_DIR/".gmail_oauth_state"
    if not state or not sp.exists() or sp.read_text(encoding="utf-8").strip()!=state: raise HTTPException(400,"OAuth state mismatch. Restart Gmail connection.")
    try:
        from google_auth_oauthlib.flow import Flow
        flow=Flow.from_client_secrets_file(str(CLIENT_SECRETS_PATH),scopes=SCOPES,state=state); flow.redirect_uri=REDIRECT_URI; flow.fetch_token(code=code)
        TOKEN_PATH.write_text(flow.credentials.to_json(),encoding="utf-8")
        try: TOKEN_PATH.chmod(0o600)
        except OSError: pass
        sp.unlink(missing_ok=True)
        return RedirectResponse(url="http://localhost:5173/?gmail_connected=1",status_code=303)
    except ImportError: raise HTTPException(400,"Install Gmail dependencies from requirements.txt first.")
    except Exception as e: raise HTTPException(400,f"OAuth callback failed: {e}") from e
