from __future__ import annotations

import io
import re
import unicodedata

try:
    import fitz  # PyMuPDF
except Exception:
    try:
        import pymupdf as fitz
    except Exception:
        fitz = None

try:
    from docx import Document
except Exception:
    Document = None


class DocumentParsingError(Exception):
    """Base exception for document parsing failures."""
    pass


class ScannedOrEmptyPdfError(DocumentParsingError):
    """Raised when a PDF contains no extractable text layer (likely scanned or image-based)."""
    pass


def clean_text(text: str) -> str:
    """Cleans, normalizes, and repairs text extracted from documents."""
    if not text:
        return ""

    # Replace null bytes
    text = text.replace("\x00", " ")

    # Unicode NFKC normalization
    text = unicodedata.normalize("NFKC", text)

    # Standardize newline representations
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Replace unicode replacement characters and non-printable control characters
    text = text.replace("\ufffd", " ")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # Standardize quotes and apostrophes
    text = re.sub(r"[\u2018\u2019\u201a\u201b`]", "'", text)
    text = re.sub(r"[\u201c\u201d\u201e\u00ab\u00bb]", '"', text)

    # Standardize dashes and hyphens
    text = re.sub(r"[\u2010\u2011\u2012\u2013\u2014\u2015\u2212]", "-", text)

    # Standardize common bullet characters into standard markdown bullets
    text = re.sub(r"^[ \t]*[\u2022\u2023\u25cf\u25cb\u25e6\u2043\u2219\u25aa\u25ab\uf0b7\uf0a7][ \t]*", "- ", text, flags=re.MULTILINE)

    # Fix broken hyphenated line wraps (e.g. "experi-\nence" -> "experience")
    text = re.sub(r"(\b[a-zA-Z]{2,})-\n([a-zA-Z]{2,}\b)", r"\1\2", text)

    # Normalize horizontal whitespace
    text = re.sub(r"[ \t]+", " ", text)

    # Remove excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def parse_pdf_bytes(raw_bytes: bytes) -> str:
    """Extracts text from PDF bytes. Detects empty/scanned PDFs."""
    if fitz is None:
        raise DocumentParsingError("PyMuPDF (fitz) is not installed. PDF parsing is unavailable.")

    try:
        doc = fitz.open(stream=raw_bytes, filetype="pdf")
    except Exception as exc:
        raise DocumentParsingError(f"Corrupted or invalid PDF file: {exc}") from exc

    page_texts = []
    for page in doc:
        text = page.get_text("text") or ""
        page_texts.append(text)

    combined = "\n".join(page_texts).strip()
    cleaned = clean_text(combined)

    # Check if PDF contains virtually no text (e.g., scanned images)
    meaningful_chars = re.sub(r"[\s\d\W_]+", "", cleaned)
    if len(meaningful_chars) < 15:
        raise ScannedOrEmptyPdfError(
            "The uploaded PDF contains no extractable text layer. "
            "It appears to be scanned or image-only. "
            "Please use OCR software to convert it, or copy/paste the text manually."
        )

    return cleaned


def parse_docx_bytes(raw_bytes: bytes) -> str:
    """Extracts text from DOCX bytes."""
    if Document is None:
        raise DocumentParsingError("python-docx is not installed. DOCX parsing is unavailable.")

    try:
        doc = Document(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise DocumentParsingError(f"Corrupted or invalid DOCX file: {exc}") from exc

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also inspect tables in DOCX
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs.append(" | ".join(row_cells))

    combined = "\n".join(paragraphs)
    cleaned = clean_text(combined)
    if not cleaned:
        raise DocumentParsingError("The DOCX document is empty.")
    return cleaned


MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB limit


def parse_document(raw_bytes: bytes, filename: str) -> str:
    """Parses an uploaded file into clean text based on its extension."""
    if not raw_bytes:
        raise DocumentParsingError("Uploaded file is empty (0 bytes).")

    if len(raw_bytes) > MAX_FILE_SIZE_BYTES:
        raise DocumentParsingError(
            f"Uploaded file exceeds maximum allowed size of 15MB ({len(raw_bytes) / (1024 * 1024):.1f}MB)."
        )

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return parse_pdf_bytes(raw_bytes)
    elif lower_name.endswith(".docx"):
        return parse_docx_bytes(raw_bytes)
    elif lower_name.endswith((".txt", ".md", ".text", ".markdown")):
        try:
            text = clean_text(raw_bytes.decode("utf-8"))
        except UnicodeDecodeError:
            try:
                text = clean_text(raw_bytes.decode("latin-1"))
            except Exception as exc:
                raise DocumentParsingError(f"Could not decode text file: {exc}") from exc
        if not text:
            raise DocumentParsingError("The text document is empty or contains only whitespace.")
        return text
    else:
        raise DocumentParsingError(f"Unsupported file format: '{filename}'. Please upload PDF, DOCX, TXT, or MD.")
